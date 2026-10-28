from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import FileResponse
import os
import glob
import time
import traceback
from datetime import datetime, timedelta, timezone
from database import db
from attendance_utils import count_present_attendance, increment_session_attendee_count
import logging
import random
import math
from typing import Optional
from pydantic import BaseModel
from pymongo.errors import DuplicateKeyError

logger = logging.getLogger(__name__)
router = APIRouter()

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "uploads")

class LocationData(BaseModel):
    latitude: float
    longitude: float

class GeofenceSession(BaseModel):
    hostId: str
    latitude: float
    longitude: float
    timeLimit: Optional[int] = None
    classTitle: Optional[str] = None
    courseCode: Optional[str] = None

class AttendanceMarkRequest(BaseModel):
    sessionId: str
    userId: str
    code: str
    latitude: float
    longitude: float
    studentName: str
    matric: str
    department: str
    level: str

class CourseAttendanceMarkRequest(BaseModel):
    userId: str
    courseCode: str
    latitude: float
    longitude: float

def generate_session_code():
    return ''.join([str(random.randint(0, 9)) for _ in range(4)])

def calculate_distance(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    # Debug logging with more precision
    logger.info(f"Calculating distance between points:")
    logger.info(f"Point 1 (Lecturer): {lat1:.6f}°N, {lon1:.6f}°E")
    logger.info(f"Point 2 (Student): {lat2:.6f}°N, {lon2:.6f}°E")

    R = 6371000  # Earth's radius in meters

    phi1 = math.radians(lat1)
    phi2 = math.radians(lat2)
    delta_phi = math.radians(lat2 - lat1)
    delta_lambda = math.radians(lon2 - lon1)

    a = math.sin(delta_phi/2) * math.sin(delta_phi/2) + \
        math.cos(phi1) * math.cos(phi2) * \
        math.sin(delta_lambda/2) * math.sin(delta_lambda/2)
    
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1-a))
    distance = R * c  # Distance in meters

    # Log with detailed precision
    logger.info(f"Calculated distance: {distance:.2f}m")
    return distance

@router.post("/start-geofence-session")
async def start_geofence_session(session_data: GeofenceSession):
    try:
        # Generate unique session ID and 4-digit code
        session_id = f"geo-{int(time.time())}"
        session_code = generate_session_code()
        
        nigeria_time = datetime.utcnow().replace(tzinfo=timezone.utc).astimezone(timezone(timedelta(hours=1)))
        
        session = {
            "sessionId": session_id,
            "sessionCode": session_code,
            "hostId": session_data.hostId,
            "classTitle": session_data.classTitle,
            "courseCode": session_data.courseCode,  # Save the course code!
            "latitude": session_data.latitude,
            "longitude": session_data.longitude,
            "radius": 100,  # Changed from 200 to 100 meters
            "status": "active",
            "isActive": True,
            "startTime": nigeria_time.isoformat(),
            "attendeeCount": 0,
            "created_at": nigeria_time.isoformat()
        }

        if session_data.timeLimit is not None:
            session["timeLimit"] = session_data.timeLimit
        
        db.sessions.insert_one(session)
        
        return {
            "status": "success",
            "sessionId": session_id,
            "sessionCode": session_code,
            "classTitle": session_data.classTitle
        }
        
    except Exception as e:
        logger.error(f"Error starting geofence session: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/mark-attendance")
def mark_attendance(request: AttendanceMarkRequest):
    try:
        # Log received coordinates
        logger.info(f"Received attendance request coordinates - Latitude: {request.latitude}, Longitude: {request.longitude}")
        
        # Get session details
        session = db.sessions.find_one({"sessionId": request.sessionId})
        if not session:
            raise HTTPException(status_code=404, detail={"status": "error", "message": "Session not found"})
            
        if session["status"] != "active":
            raise HTTPException(status_code=400, detail={"status": "error", "message": "Session has ended"})
            
        if session["sessionCode"] != request.code:
            raise HTTPException(status_code=400, detail={"status": "error", "message": "Invalid session code"})
        
        # Log session coordinates for comparison
        logger.info(f"Session location - Latitude: {session['latitude']}, Longitude: {session['longitude']}")
        
        # Calculate distance between lecturer and student
        distance = calculate_distance(
            session["latitude"], 
            session["longitude"],
            request.latitude,
            request.longitude
        )
        
        # Log calculated distance
        logger.info(f"Calculated distance: {int(distance)}m")
        
        if distance > session["radius"]:
            raise HTTPException(
                status_code=400, 
                detail={"status": "error", "message": f"You are {int(distance)}m away from the lecture location. Must be within {session['radius']}m"}
            )
            
        # Record attendance
        nigeria_time = datetime.utcnow().replace(tzinfo=timezone.utc).astimezone(timezone(timedelta(hours=1)))
        attendance_record = {
            "sessionId": request.sessionId,
            "userId": request.userId,
            "attendanceStatus": "present",
            "name": request.studentName,
            "matric": request.matric,
            "department": request.department,
            "level": request.level,
            "timestamp": nigeria_time.isoformat(),
            "distance": int(distance),
            "latitude": request.latitude,
            "longitude": request.longitude
        }
        
        try:
            db.attendance.insert_one(attendance_record)
        except DuplicateKeyError:
            raise HTTPException(status_code=400, detail={"status": "error", "message": "Attendance already marked"})

        increment_session_attendee_count(db, request.sessionId)
        
        return {
            "status": "success",
            "message": "Attendance marked successfully",
            "data": {
                "timestamp": nigeria_time.isoformat(),
                "distance": int(distance)
            }
        }
        
    except HTTPException as he:
        raise he
    except Exception as e:
        logger.error(f"Error marking attendance: {str(e)}")
        raise HTTPException(
            status_code=500, 
            detail={"status": "error", "message": str(e)}
        )

@router.get("/download-attendance")
async def download_attendance(sessionId: str, userId: str):
    if not sessionId or not userId:
        logger.warning(f"Missing parameters - sessionId: {sessionId}, userId: {userId}")
        raise HTTPException(status_code=400, detail="Missing required parameters")

    try:
        logger.info(f"Attendance download requested - Session: {sessionId}, User: {userId}")
        
        # First check database for session and attendance records
        session = db.sessions.find_one({"sessionId": sessionId})
        if not session:
            logger.warning(f"Session not found: {sessionId}")
            raise HTTPException(status_code=404, detail="Session not found")
            
        # Check user authorization
        host_id = session.get("hostId")
        if not host_id:
            logger.error(f"Session {sessionId} has no host ID")
            raise HTTPException(status_code=500, detail="Invalid session configuration")
            
        if host_id != userId:
            logger.warning(f"Unauthorized access attempt by user {userId} for session {sessionId}")
            raise HTTPException(status_code=403, detail="Unauthorized access")

        # Ensure session is completed
        if session.get("status") != "completed":
            logger.warning(f"Attempted to download active session: {sessionId}")
            raise HTTPException(
                status_code=400, 
                detail="Cannot download attendance for active session. Please end the session first."
            )

        # Get attendance records from database
        records = list(db.attendance.find(
            {"sessionId": sessionId},
            {"_id": 0}
        ).sort("created_at", 1))
        present_count = count_present_attendance(db, sessionId)

        if not records:
            logger.warning(f"No attendance records found for session {sessionId}")
            raise HTTPException(
                status_code=404,
                detail="No attendance records found for this session"
            )

        # Get lecturer information
        lecturer = db.users.find_one({"userId": host_id})
        lecturer_name = lecturer.get("name", "Unknown Lecturer") if lecturer else "Unknown Lecturer"
        
        # Get course code from session
        course_code = session.get("courseCode", "N/A")
        
        # Get session geofence radius
        geofence_radius = session.get("radius", 100)

        # Generate Excel file
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"

        # Add session information
        ws.merge_cells('A1:F1')
        course_title = session.get('classTitle', 'Untitled Session')
        ws['A1'] = f"Attendance Records: {course_title}"
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

        # Add metadata
        ws['A2'] = f"Course Code: {course_code}"
        ws['A3'] = f"Lecturer: {lecturer_name}"
        ws['A4'] = f"Start Time: {session.get('startTime', 'N/A')}"
        ws['A5'] = f"End Time: {session.get('endTime', 'N/A')}"
        ws['A6'] = f"Total Present: {present_count}"
        ws['A7'] = f"Generated at: {datetime.now(timezone(timedelta(hours=1))).strftime('%Y-%m-%d %H:%M:%S %Z')}"

        # Add headers
        headers = ["Name", "Matric Number", "Department", "Time", "Status", "Verification Score"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=8, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

        # Add records
        for row, record in enumerate(records, 9):
            ws.cell(row=row, column=1).value = record.get("name", "")
            ws.cell(row=row, column=2).value = record.get("matric", "")
            ws.cell(row=row, column=3).value = record.get("department", "")

            # Format timestamp
            try:
                if not record.get("timestamp"):
                    raise ValueError("No attendance timestamp")
                timestamp = datetime.fromisoformat(record.get("timestamp", ""))
                if not timestamp.tzinfo:
                    timestamp = timestamp.replace(tzinfo=timezone.utc)
                nigeria_time = timestamp.astimezone(timezone(timedelta(hours=1)))
                ws.cell(row=row, column=4).value = nigeria_time.strftime("%Y-%m-%d %I:%M %p")
            except Exception as e:
                logger.error(f"Error formatting timestamp: {e}")
                ws.cell(row=row, column=4).value = record.get("timestamp", "")

            attendance_status = record.get("attendanceStatus", "present")
            ws.cell(row=row, column=5).value = attendance_status.title()
            
            # Calculate verification score based on geofencing accuracy
            # Verification score = how accurately the student was within the geofence
            # 100% = at lecture location (0m), decreasing as they get closer to boundary
            student_distance = record.get("distance", geofence_radius)
            if attendance_status == "present" and student_distance is not None and geofence_radius > 0:
                # If student is at the center, score is 100%
                # If student is at the edge of geofence, score is lower
                # Formula: (1 - (distance / radius)) * 100
                verification_score = max(0, (1 - (float(student_distance) / float(geofence_radius))) * 100)
                ws.cell(row=row, column=6).value = f"{round(verification_score, 2)}%"
            else:
                ws.cell(row=row, column=6).value = "100%"

        # Adjust column widths
        for col in range(1, 7):
            ws.column_dimensions[chr(64 + col)].width = 15

        # Save the file
        if not os.path.exists(UPLOAD_DIR):
            os.makedirs(UPLOAD_DIR)

        filename = f"attendance_{sessionId}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        file_path = os.path.join(UPLOAD_DIR, filename)
        
        wb.save(file_path)
        
        # Verify file was created and has content
        if not os.path.exists(file_path):
            raise HTTPException(status_code=500, detail="Failed to generate attendance file")
            
        file_size = os.path.getsize(file_path)
        if file_size == 0:
            os.remove(file_path)
            raise HTTPException(status_code=500, detail="Generated file is empty")

        return FileResponse(
            file_path,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=filename,
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )

    except HTTPException as http_err:
        raise http_err
    except Exception as e:
        logger.error(f"Error generating attendance report: {str(e)}")
        return HTTPException(status_code=500, detail=f"Failed to generate attendance report: {str(e)}")

@router.post("/end-attendance-session")
async def end_attendance_session(request: Request):
    try:
        data = await request.json()
        session_id = data.get("sessionId")
        user_id = data.get("userId")
        
        logger.info(f"[end-attendance-session] Received request: sessionId={session_id}, userId={user_id}")
        
        if not session_id or not user_id:
            logger.error(f"[end-attendance-session] Missing required fields")
            raise HTTPException(status_code=400, detail="Missing required fields")

        # Get session info
        session = db.sessions.find_one({"sessionId": session_id})
        logger.info(f"[end-attendance-session] Found session: {session}")
        if not session:
            logger.error(f"[end-attendance-session] Session not found: {session_id}")
            raise HTTPException(status_code=404, detail="Session not found")
            
        if session.get("hostId") != user_id:
            logger.error(f"[end-attendance-session] Unauthorized: hostId={session.get('hostId')}, userId={user_id}")
            raise HTTPException(status_code=403, detail="Unauthorized to end this session")
        
        current_status = session.get("status")
        logger.info(f"[end-attendance-session] Current session status: {current_status}")
            
        # If session is already completed, just return success with existing data
        if current_status == "completed":
            logger.info(f"[end-attendance-session] Session already completed: {session_id}")
            finalization_status = session.get("absenceFinalizationStatus", "completed")
            if "absenceFinalizationStatus" not in session:
                db.sessions.update_one(
                    {"sessionId": session_id, "absenceFinalizationStatus": {"$exists": False}},
                    {"$set": {
                        "absenceFinalizationStatus": "pending",
                        "absenceFinalizationQueuedAt": datetime.utcnow(),
                    }},
                )
                finalization_status = "pending"
            return {
                "status": "success",
                "message": "Session was already completed",
                "recordCount": count_present_attendance(db, session_id),
                "absenceFinalizationStatus": finalization_status,
                "sessionEndTime": session.get("endTime", "N/A"),
                "alreadyEnded": True
            }
            
        # If status is not active, log it but continue anyway (it might have been ended in another way)
        if current_status != "active":
            logger.warning(f"[end-attendance-session] Session status is {current_status}, but will attempt to end it anyway")

        # Get all records and archive them
        records = list(db.attendance.find({"sessionId": session_id}).sort("created_at", 1))
        logger.info(f"[end-attendance-session] Found {len(records)} attendance records")
        
        nigeria_time = datetime.utcnow().replace(tzinfo=timezone.utc).astimezone(timezone(timedelta(hours=1)))
        
        # Update each attendance record with session end details
        for record in records:
            logger.info(f"[end-attendance-session] Updating attendance record: {record['_id']}")
            db.attendance.update_one(
                {"_id": record["_id"]},
                {
                    "$set": {
                        "sessionEndTime": nigeria_time.isoformat(),
                        "sessionStatus": "completed",
                        "lastUpdated": datetime.utcnow().isoformat()
                    }
                }
            )
        
        # Update session status
        logger.info(f"[end-attendance-session] Updating session status to completed")
        result = db.sessions.update_one(
            {"sessionId": session_id},
            {
                "$set": {
                    "status": "completed",
                    "endTime": nigeria_time.isoformat(),
                    "attendeeCount": count_present_attendance(db, session_id),
                    "absenceFinalizationStatus": "pending",
                    "absenceFinalizationQueuedAt": nigeria_time,
                    "lastUpdated": datetime.utcnow().isoformat()
                }
            }
        )

        logger.info(f"[end-attendance-session] Session update result: matched={result.matched_count}, modified={result.modified_count}")
        
        # If this session was created from a schedule, update the schedule status too
        if session.get("scheduleId"):
            logger.info(f"[end-attendance-session] Updating schedule status for {session.get('scheduleId')}")
            db.schedules.update_one(
                {"scheduleId": session.get("scheduleId")},
                {
                    "$set": {
                        "status": "completed",
                        "updatedAt": nigeria_time.isoformat()
                    }
                }
            )
            logger.info(f"[end-attendance-session] Schedule {session.get('scheduleId')} marked as completed")
        
        return {
            "status": "success",
            "message": "Attendance session ended successfully",
            "downloadUrl": f"/download-attendance/{session_id}",
            "recordCount": count_present_attendance(db, session_id),
            "absenceFinalizationStatus": "pending",
            "sessionEndTime": nigeria_time.isoformat()
        }
        
    except HTTPException as he:
        logger.error(f"[end-attendance-session] HTTPException: {he.detail}")
        raise he
    except Exception as e:
        logger.error(f"[end-attendance-session] Unexpected error: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/get-active-session")
async def get_active_session(code: str):
    try:
        # Find active session with this code
        session = db.sessions.find_one({
            "sessionCode": code,
            "status": "active"
        })
        
        if not session:
            raise HTTPException(status_code=404, detail="No active session found with this code")
        
        return {
            "status": "success",
            "sessionId": session["sessionId"],
            "message": "Active session found"
        }
        
    except HTTPException as he:
        raise he
    except Exception as e:
        logger.error(f"Error finding active session: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/get-student-details")
async def get_student_details(userId: str):
    try:
        # Get user details from database
        user = db.users.find_one({"userId": userId})
        if not user:
            raise HTTPException(status_code=404, detail={"status": "error", "message": "Student not found"})
            
        if user.get("role") != "student":
            raise HTTPException(status_code=400, detail={"status": "error", "message": "User is not a student"})

        return {
            "status": "success",
            "data": {
                "name": user.get("name", ""),
                "matricNumber": user.get("matricNumber", ""),
                "department": user.get("department", ""),
                "level": user.get("level", "")
            }
        }
    except HTTPException as he:
        raise he
    except Exception as e:
        logger.error(f"Error getting student details: {str(e)}")
        raise HTTPException(
            status_code=500,
            detail={"status": "error", "message": str(e)}
        )

@router.get("/get-session-status")
async def get_session_status(sessionId: str):
    try:
        session = db.sessions.find_one({"sessionId": sessionId})
        if not session:
            raise HTTPException(status_code=404, detail="Session not found")

        # Calculate session duration and check if it's still valid
        start_time = datetime.fromisoformat(session["startTime"].replace('Z', '+00:00'))
        current_time = datetime.utcnow().replace(tzinfo=timezone.utc)
        session_duration = (current_time - start_time).total_seconds() / 3600  # in hours

        # Update session status if it's been active for too long (e.g., 24 hours)
        if session["status"] == "active" and session_duration > 24:
            db.sessions.update_one(
                {"sessionId": sessionId},
                {
                    "$set": {
                        "status": "completed",
                        "endTime": current_time.isoformat(),
                        "absenceFinalizationStatus": "pending",
                        "absenceFinalizationQueuedAt": current_time,
                        "lastUpdated": current_time.isoformat()
                    }
                }
            )
            return {
                "status": "completed",
                "isActive": False,
                "message": "Session automatically ended due to duration limit"
            }

        return {
            "status": session["status"],
            "isActive": session["status"] == "active",
            "startTime": session["startTime"],
            "attendeeCount": session.get("attendeeCount", 0),
            "message": "Session status retrieved successfully"
        }
    except HTTPException as he:
        raise he
    except Exception as e:
        logger.error(f"Error getting session status: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/get-lecturer-active-session")
async def get_lecturer_active_session(userId: str):
    """Get active session for a lecturer if one exists"""
    try:
        # Find active session for this lecturer
        active_session = db.sessions.find_one({
            "hostId": userId,
            "status": "active"
        })
        
        if not active_session:
            return {"session": None}
        
        logger.info(f"Found active session for lecturer {userId}: {active_session.get('sessionId')}")
        
        return {
            "session": {
                "sessionId": active_session.get("sessionId"),
                "courseCode": active_session.get("courseCode"),
                "startTime": active_session.get("startTime"),
                "attendeeCount": active_session.get("attendeeCount", 0)
            }
        }
    except Exception as e:
        logger.error(f"Error getting lecturer active session: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/check-attendance-status")
async def check_attendance_status(userId: str, sessionId: str):
    try:
        # Check if session exists and is active
        session = db.sessions.find_one({
            "sessionId": sessionId,
            "status": "active"
        })
        
        if not session:
            return {
                "status": "error",
                "hasMarkedAttendance": False,
                "message": "Session not found or not active"
            }

        # Check if student has marked attendance
        attendance = db.attendance.find_one({
            "sessionId": sessionId,
            "userId": userId
        })

        return {
            "status": "success",
            "hasMarkedAttendance": bool(attendance),
            "message": "Attendance already marked" if attendance else "Attendance not yet marked",
            "timestamp": attendance["timestamp"] if attendance else None
        }
    except Exception as e:
        logger.error(f"Error checking attendance status: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/get-session-course-code")
async def get_session_course_code(sessionId: str):
    """Get the course code for a geofence session"""
    try:
        # Find the session
        session = db.sessions.find_one({
            "sessionId": sessionId
        })
        
        if not session:
            return {
                "status": "error",
                "message": "Session not found",
                "code": None,
                "courseCode": None
            }
        
        # Return the course code from the session
        courseCode = session.get("courseCode") or session.get("code")
        
        return {
            "status": "success",
            "code": courseCode,
            "courseCode": courseCode,
            "sessionId": sessionId
        }
    except Exception as e:
        logger.error(f"Error getting session course code: {str(e)}")
        return {
            "status": "error",
            "message": str(e),
            "code": None,
            "courseCode": None
        }

# Test coordinates
LECTURER_TEST_COORDS = {
    "latitude": 6.827302,
    "longitude": 3.940243
}

STUDENT_TEST_COORDS = {
    "latitude": 6.827402,
    "longitude": 3.940343
}

# Add this function to test coordinates
def test_distance():
    distance = calculate_distance(
        LECTURER_TEST_COORDS["latitude"],
        LECTURER_TEST_COORDS["longitude"],
        STUDENT_TEST_COORDS["latitude"],
        STUDENT_TEST_COORDS["longitude"]
    )
    logger.info(f"Test distance between points: {distance:.2f}m")
    return distance <= 100  # Changed from 200 to 100

@router.get("/test-coords")
async def test_coordinates():
    distance = calculate_distance(
        LECTURER_TEST_COORDS["latitude"],
        LECTURER_TEST_COORDS["longitude"],
        STUDENT_TEST_COORDS["latitude"],
        STUDENT_TEST_COORDS["longitude"]
    )
    return {
        "distance": round(distance, 2),
        "within_range": distance <= 100,  # Changed from 200 to 100
        "lecturer": LECTURER_TEST_COORDS,
        "student": STUDENT_TEST_COORDS
    }

@router.get("/get-student-courses")
async def get_student_courses(userId: str):
    """
    Get all courses the student registered for, with active session status.
    
    This endpoint returns all registered courses and marks which ones have active attendance sessions.
    Active sessions can be created in two ways:
    1. Manually started by lecturer (via /start-geofence-session)
    2. Automatically started from a scheduled session (via /start-scheduled-session in schedule.py)
    
    Both types of sessions are stored in the same 'sessions' collection with:
    - status: "active"
    - courseCode: matching the student's registered course code
    - sessionId and hostId for the student to join
    
    Students will see scheduled courses as "active" the same way as manually started ones.
    """
    try:
        from bson import ObjectId
        
        logger.info(f"[get-student-courses] Starting for userId: {userId}")
        
        # Get student details
        student = db.users.find_one({"userId": userId, "role": "student"})
        if not student:
            logger.warning(f"[get-student-courses] Student not found for userId: {userId}")
            raise HTTPException(status_code=404, detail="Student not found")
        
        logger.info(f"[get-student-courses] Found student: {student.get('name')}")
        
        # Get student's registered courses (course IDs/codes from registration)
        registered_course_ids = student.get("courses", [])
        logger.info(f"[get-student-courses] Student's registered course IDs: {registered_course_ids}")
        
        if not registered_course_ids:
            logger.info(f"[get-student-courses] Student has no registered courses")
            return {
                "status": "success",
                "courses": [],
                "message": "You haven't registered for any courses yet"
            }
        
        # Get course details for each registered course
        student_courses = []
        try:
            for course_id in registered_course_ids:
                logger.info(f"[get-student-courses] Processing course ID: {course_id}")
                # Try to fetch as ObjectId first
                course = None
                try:
                    if isinstance(course_id, str) and len(course_id) == 24:  # ObjectId length
                        obj_id = ObjectId(course_id)
                        course = db.courses.find_one({"_id": obj_id})
                        if course:
                            logger.info(f"[get-student-courses] Found course by ObjectId {course_id}: {course.get('code')}")
                except Exception as e:
                    logger.warning(f"[get-student-courses] Failed to find by ObjectId: {str(e)}")
                
                # If ObjectId fails, try searching by code
                if not course:
                    course = db.courses.find_one({"code": course_id})
                    if course:
                        logger.info(f"[get-student-courses] Found course by code {course_id}: {course.get('code')}")
                
                if course:
                    course_data = {
                        "code": course.get("code", ""),
                        "title": course.get("title", ""),
                        "_id": course_id  # Keep the original ID for reference
                    }
                    student_courses.append(course_data)
                    logger.info(f"[get-student-courses] Added course to response: {course_data['code']}")
                else:
                    logger.warning(f"[get-student-courses] Course not found for ID: {course_id}")
        except Exception as e:
            logger.error(f"[get-student-courses] Error fetching course details: {str(e)}")
            logger.error(traceback.format_exc())
        
        logger.info(f"[get-student-courses] Loaded {len(student_courses)} courses")
        
        if not student_courses:
            logger.warning(f"[get-student-courses] No course details found despite having course IDs")
            return {
                "status": "success",
                "courses": [],
                "message": "No course details found for your registered courses"
            }
        
        # Get course codes
        course_codes = [course.get("code") for course in student_courses]
        logger.info(f"[get-student-courses] Course codes to search: {course_codes}")
        
        # Find all active sessions with matching course codes
        active_sessions = list(db.sessions.find({
            "courseCode": {"$in": course_codes},
            "status": "active"
        }))
        
        logger.info(f"[get-student-courses] Query: courseCode in {course_codes} AND status='active'")
        logger.info(f"[get-student-courses] Found {len(active_sessions)} active sessions")
        
        # Debug: Log all sessions with matching course codes (regardless of status)
        all_matching_sessions = list(db.sessions.find({
            "courseCode": {"$in": course_codes}
        }))
        logger.info(f"[get-student-courses] Total sessions with matching course codes (all statuses): {len(all_matching_sessions)}")
        for sess in all_matching_sessions:
            logger.info(f"[get-student-courses] Session: {sess.get('sessionId')} | Course: {sess.get('courseCode')} | Status: {sess.get('status')}")
        
        # Return all registered courses with active session info
        courses_response = []
        for course in student_courses:
            course_code = course.get("code", "")
            # Check if there's an active session for this course
            session = next((s for s in active_sessions if s.get("courseCode") == course_code), None)
            
            course_data = {
                "code": course_code,
                "title": course.get("title", ""),
                "hasActiveSession": bool(session),
            }
            
            if session:
                course_data["sessionId"] = session.get("sessionId")
                course_data["hostId"] = session.get("hostId")
                logger.info(f"[get-student-courses] Course {course_code} has active session: {session.get('sessionId')}")
            else:
                logger.info(f"[get-student-courses] Course {course_code} has NO active session")
            
            courses_response.append(course_data)
        
        logger.info(f"[get-student-courses] Returning {len(courses_response)} courses for student {userId}")
        
        return {
            "status": "success",
            "courses": courses_response,
            "count": len(courses_response)
        }
        
    except HTTPException as he:
        raise he
    except Exception as e:
        logger.error(f"[get-student-courses] Error: {str(e)}")
        logger.error(traceback.format_exc())
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/mark-attendance-by-course")
def mark_attendance_by_course(request: CourseAttendanceMarkRequest):
    """Mark attendance by selecting a course (replaces code entry)"""
    try:
        user_id = request.userId
        course_code = request.courseCode
        latitude = request.latitude
        longitude = request.longitude
        
        if not all([user_id, course_code, latitude is not None, longitude is not None]):
            raise HTTPException(status_code=400, detail="Missing required fields")
        
        # Get student details
        student = db.users.find_one({"userId": user_id, "role": "student"})
        if not student:
            raise HTTPException(status_code=404, detail="Student not found")
        
        # Check if student is restricted
        if student.get("isRestricted", False):
            raise HTTPException(status_code=403, detail="Your account has been restricted. Please contact an administrator.")
        
        # Find active session for this course
        session = db.sessions.find_one({
            "courseCode": course_code,
            "status": "active"
        })
        
        if not session:
            raise HTTPException(status_code=404, detail="No active session for this course")
        
        session_id = session.get("sessionId")
        
        # Verify geolocation (student must be within radius)
        distance = calculate_distance(
            session["latitude"],
            session["longitude"],
            latitude,
            longitude
        )
        
        if distance > session.get("radius", 100):
            raise HTTPException(
                status_code=400,
                detail=f"You are {int(distance)}m away from the lecture location. Must be within {session.get('radius', 100)}m"
            )
        
        # Record attendance
        nigeria_time = datetime.utcnow().replace(tzinfo=timezone.utc).astimezone(timezone(timedelta(hours=1)))
        attendance_record = {
            "sessionId": session_id,
            "userId": user_id,
            "attendanceStatus": "present",
            "courseCode": course_code,
            "name": student.get("name", ""),
            "matric": student.get("matric", ""),
            "department": student.get("department", ""),
            "level": student.get("level", ""),
            "timestamp": nigeria_time.isoformat(),
            "distance": int(distance),
            "latitude": latitude,
            "longitude": longitude,
            "created_at": nigeria_time.isoformat()
        }
        
        try:
            db.attendance.insert_one(attendance_record)
        except DuplicateKeyError:
            raise HTTPException(status_code=400, detail="Attendance already marked for this session")

        increment_session_attendee_count(db, session_id)
        
        return {
            "status": "success",
            "message": "Attendance marked successfully",
            "data": {
                "courseCode": course_code,
                "timestamp": nigeria_time.isoformat(),
                "distance": int(distance)
            }
        }
        
    except HTTPException as he:
        raise he
    except Exception as e:
        logger.error(f"Error marking attendance by course: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/get-session-attendance")
async def get_session_attendance(sessionId: str):
    """
    Get real-time attendance records for a session
    """
    try:
        session = db.sessions.find_one({"sessionId": sessionId})
        if not session:
            raise HTTPException(status_code=404, detail="Session not found")
        
        # Fetch all attendance records for this session, sorted by timestamp (newest first)
        attendance_records = list(db.attendance.find(
            {"sessionId": sessionId}
        ).sort("timestamp", -1).limit(100))  # Limit to last 100 records for performance
        
        # Format attendance data for frontend
        formatted_attendance = []
        for record in attendance_records:
            formatted_attendance.append({
                "studentId": record.get("userId", ""),
                "studentName": record.get("name", "Unknown Student"),
                "timestamp": record.get("timestamp", ""),
                "distance": record.get("distance", None),
                "status": record.get("attendanceStatus", "present")
            })
        
        return {
            "status": "success",
            "attendance": formatted_attendance,
            "totalCount": count_present_attendance(db, sessionId),
            "absentCount": sum(1 for record in formatted_attendance if record["status"] == "absent")
        }
        
    except HTTPException as he:
        raise he
    except Exception as e:
        logger.error(f"Error fetching session attendance: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))


def _attendance_report_context(user_id):
    from bson import ObjectId

    lecturer = db.users.find_one({"userId": user_id})
    if not lecturer:
        raise HTTPException(status_code=404, detail="Lecturer not found")
    if lecturer.get("role") != "lecturer":
        raise HTTPException(status_code=403, detail="Only lecturers can view attendance reports")

    sessions = list(db.sessions.find({"hostId": user_id}, {"_id": 0}).sort("created_at", -1))
    course_refs = lecturer.get("courses", [])
    course_ids = []
    course_codes = set()
    for course_ref in course_refs:
        if isinstance(course_ref, ObjectId):
            course_ids.append(course_ref)
        elif isinstance(course_ref, str):
            course_codes.add(course_ref)
            if ObjectId.is_valid(course_ref):
                course_ids.append(ObjectId(course_ref))

    course_codes.update(
        code for session in sessions
        if (code := session.get("courseCode") or session.get("code"))
    )
    course_query = []
    if course_ids:
        course_query.append({"_id": {"$in": course_ids}})
    if course_codes:
        course_query.append({"code": {"$in": list(course_codes)}})

    course_lookup = {}
    if course_query:
        for course in db.courses.find(
            {"$or": course_query},
            {"_id": 1, "code": 1, "title": 1},
        ):
            if course.get("code"):
                course_lookup[course["code"]] = course

    for session in sessions:
        code = session.get("courseCode") or session.get("code")
        if code and code not in course_lookup:
            course_lookup[code] = {
                "code": code,
                "title": session.get("classTitle") or "Untitled course",
            }

    course_docs = [
        {"code": code, "title": course.get("title", "Untitled course")}
        for code, course in course_lookup.items()
    ]
    return lecturer, sessions, course_lookup, course_docs


def _course_enrollment_refs(course_lookup):
    refs_by_code = {}
    codes_by_ref = {}
    for code, course in course_lookup.items():
        refs = {code, str(course.get("_id", ""))}
        if course.get("_id"):
            refs.add(course["_id"])
        refs_by_code[code] = refs
        for ref in refs:
            codes_by_ref.setdefault(str(ref), set()).add(code)
    return refs_by_code, codes_by_ref


def _session_summary(session, course_lookup, present_count, enrolled_count):
    session_id = session.get("sessionId", "")
    course_code = session.get("courseCode") or session.get("code") or ""
    course = course_lookup.get(course_code, {})
    return {
        "sessionId": session_id,
        "courseCode": course_code,
        "courseTitle": course.get("title") or session.get("classTitle") or "Untitled session",
        "type": "physical" if session_id.startswith("geo-") else "virtual",
        "status": session.get("status", "completed"),
        "finalizationStatus": session.get("absenceFinalizationStatus", "completed"),
        "startTime": session.get("startTime", ""),
        "endTime": session.get("endTime", ""),
        "createdAt": session.get("created_at") or session.get("createdAt") or "",
        "attendeeCount": present_count,
        "enrolledCount": enrolled_count,
        "attendanceRate": round(present_count / enrolled_count * 100, 1) if enrolled_count else None,
    }


class AttendanceReportDetailsRequest(BaseModel):
    userId: str
    sessionIds: list[str]


def _attendance_report_details(sessions, course_lookup):
    session_ids = [session.get("sessionId") for session in sessions if session.get("sessionId")]
    attendance_rows = list(db.attendance.find(
        {"sessionId": {"$in": session_ids}},
        {"_id": 0, "sessionId": 1, "userId": 1, "name": 1, "matric": 1,
         "attendanceStatus": 1, "department": 1, "level": 1, "timestamp": 1},
    )) if session_ids else []

    refs_by_code, codes_by_ref = _course_enrollment_refs(course_lookup)
    all_refs = [ref for refs in refs_by_code.values() for ref in refs]
    roster_rows = list(db.users.find(
        {"role": "student", "courses": {"$in": all_refs}},
        {"_id": 0, "userId": 1, "name": 1, "matricNumber": 1, "matric": 1,
         "department": 1, "level": 1, "courses": 1},
    )) if all_refs else []

    attendance_by_session = {}
    students_by_id = {}
    for row in attendance_rows:
        session_id = row.get("sessionId")
        student_id = row.get("userId") or row.get("matric") or row.get("name")
        row["attendanceStatus"] = row.get("attendanceStatus", "present")
        attendance_by_session.setdefault(session_id, {})[student_id] = row
        students_by_id[student_id] = {
            "userId": row.get("userId", ""),
            "name": row.get("name", "Unknown student"),
            "matric": row.get("matric", ""),
            "department": row.get("department", ""),
            "level": row.get("level", ""),
        }

    roster_by_code = {code: {} for code in course_lookup}
    for student in roster_rows:
        student_id = student.get("userId") or student.get("matricNumber") or student.get("matric")
        matched_codes = {
            code
            for ref in student.get("courses", [])
            for code in codes_by_ref.get(str(ref), ())
        }
        for code in matched_codes:
            if student_id:
                roster_by_code[code][student_id] = student
        if student_id:
            students_by_id.setdefault(student_id, {
                "userId": student.get("userId", ""),
                "name": student.get("name", "Unknown student"),
                "matric": student.get("matricNumber") or student.get("matric", ""),
                "department": student.get("department", ""),
                "level": student.get("level", ""),
            })

    details_by_session = {}
    for session in sessions:
        session_id = session.get("sessionId", "")
        course_code = session.get("courseCode") or session.get("code") or ""
        present_by_id = attendance_by_session.get(session_id, {})
        roster = roster_by_code.get(course_code, {})
        attendees = [{
            "userId": row.get("userId", ""),
            "name": row.get("name", "Unknown student"),
            "matric": row.get("matric", ""),
            "department": row.get("department", ""),
            "level": row.get("level", ""),
            "timestamp": row.get("timestamp", ""),
            "status": row.get("attendanceStatus", "present"),
        } for row in present_by_id.values()]
        for student_id, student in roster.items():
            if student_id not in present_by_id:
                attendees.append({
                    "userId": student.get("userId", ""),
                    "name": student.get("name", "Unknown student"),
                    "matric": student.get("matricNumber") or student.get("matric", ""),
                    "department": student.get("department", ""),
                    "level": student.get("level", ""),
                    "timestamp": "",
                    "status": "absent",
                })
        attendees.sort(key=lambda attendee: (attendee["status"] != "present", attendee["name"].lower()))
        details_by_session[session_id] = {
            **_session_summary(
                session,
                course_lookup,
                sum(1 for row in present_by_id.values() if row.get("attendanceStatus", "present") == "present"),
                len(roster) if course_code in roster_by_code else None,
            ),
            "attendees": attendees,
        }
    return details_by_session, list(students_by_id.values())


@router.get("/attendance-report")
async def get_attendance_report(userId: str):
    """Return lightweight attendance summaries for a lecturer."""
    try:
        _, sessions, course_lookup, course_docs = _attendance_report_context(userId)
        session_ids = [session.get("sessionId") for session in sessions if session.get("sessionId")]
        attendance_counts = {}
        if session_ids:
            for item in db.attendance.aggregate([
                {"$match": {"sessionId": {"$in": session_ids}}},
                {"$group": {
                    "_id": "$sessionId",
                    "presentCount": {"$sum": {"$cond": [
                        {"$ne": [{"$ifNull": ["$attendanceStatus", "present"]}, "absent"]},
                        1,
                        0,
                    ]}},
                }},
            ]):
                attendance_counts[item["_id"]] = item.get("presentCount", 0)

        refs_by_code, codes_by_ref = _course_enrollment_refs(course_lookup)
        all_refs = [ref for refs in refs_by_code.values() for ref in refs]
        roster_counts = {code: 0 for code in course_lookup}
        if all_refs:
            roster_projection = {
                "_id": 0, "userId": 1, "matricNumber": 1, "matric": 1, "courses": 1,
            }
            for student in db.users.find(
                {"role": "student", "courses": {"$in": all_refs}},
                roster_projection,
            ):
                student_id = student.get("userId") or student.get("matricNumber") or student.get("matric")
                if not student_id:
                    continue
                for ref in {str(ref) for ref in student.get("courses", [])}:
                    for code in codes_by_ref.get(ref, ()):
                        roster_counts[code] += 1

        summaries = []
        for session in sessions:
            session_id = session.get("sessionId", "")
            code = session.get("courseCode") or session.get("code") or ""
            present_count = attendance_counts.get(session_id, 0)
            summaries.append(_session_summary(
                session,
                course_lookup,
                present_count,
                roster_counts.get(code) if code in roster_counts else None,
            ))
        return {"courses": course_docs, "sessions": summaries}
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Error building lecturer attendance summary")
        raise HTTPException(status_code=500, detail="Failed to build attendance report") from e


@router.get("/attendance-report/session/{session_id}")
async def get_attendance_report_session(session_id: str, userId: str):
    try:
        _, sessions, course_lookup, _ = _attendance_report_context(userId)
        session = next((item for item in sessions if item.get("sessionId") == session_id), None)
        if not session:
            raise HTTPException(status_code=404, detail="Session not found")
        code = session.get("courseCode") or session.get("code") or ""
        details, _ = _attendance_report_details(
            [session],
            {code: course_lookup[code]} if code in course_lookup else {},
        )
        return details[session_id]
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Error loading attendance session details")
        raise HTTPException(status_code=500, detail="Failed to load attendance session details") from e


@router.post("/attendance-report/session-details")
async def get_attendance_report_session_details(request: AttendanceReportDetailsRequest):
    try:
        _, sessions, course_lookup, _ = _attendance_report_context(request.userId)
        requested_ids = set(request.sessionIds)
        selected_sessions = [
            session for session in sessions
            if session.get("sessionId") in requested_ids
        ]
        selected_codes = {
            session.get("courseCode") or session.get("code")
            for session in selected_sessions
            if session.get("courseCode") or session.get("code")
        }
        scoped_courses = {
            code: course_lookup[code]
            for code in selected_codes
            if code in course_lookup
        }
        details, _ = _attendance_report_details(selected_sessions, scoped_courses)
        return {"sessions": list(details.values())}
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Error loading attendance export details")
        raise HTTPException(status_code=500, detail="Failed to load attendance export details") from e


@router.get("/attendance-report/students")
async def get_attendance_report_students(userId: str):
    try:
        _, sessions, course_lookup, _ = _attendance_report_context(userId)
        _, students = _attendance_report_details([], course_lookup)
        session_ids = [session.get("sessionId") for session in sessions if session.get("sessionId")]
        attendance_rows = list(db.attendance.find(
            {"sessionId": {"$in": session_ids}},
            {"_id": 0, "userId": 1, "name": 1, "matric": 1, "department": 1, "level": 1},
        )) if session_ids else []
        students_by_key = {
            student.get("userId") or student.get("matric") or student.get("name"): student
            for student in students
        }
        for row in attendance_rows:
            key = row.get("userId") or row.get("matric") or row.get("name")
            if key:
                students_by_key.setdefault(key, {
                    "userId": row.get("userId", ""),
                    "name": row.get("name", "Unknown student"),
                    "matric": row.get("matric", ""),
                    "department": row.get("department", ""),
                    "level": row.get("level", ""),
                })
        return {"students": sorted(students_by_key.values(), key=lambda student: student["name"].lower())}
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Error loading attendance report students")
        raise HTTPException(status_code=500, detail="Failed to load report students") from e


@router.get("/attendance-report/student-history")
async def get_attendance_report_student_history(userId: str, studentId: str):
    try:
        _, sessions, course_lookup, _ = _attendance_report_context(userId)
        student = db.users.find_one({
            "role": "student",
            "$or": [{"userId": studentId}, {"matricNumber": studentId}, {"matric": studentId}],
        })
        refs_by_code, codes_by_ref = _course_enrollment_refs(course_lookup)
        enrollment_refs = {str(ref) for ref in (student or {}).get("courses", [])}
        enrolled_codes = {
            code
            for ref in enrollment_refs
            for code in codes_by_ref.get(ref, ())
        }
        session_ids = [session.get("sessionId") for session in sessions if session.get("sessionId")]
        attendance_filter = [{"userId": (student or {}).get("userId", studentId)}]
        if (student or {}).get("matricNumber") or (student or {}).get("matric"):
            attendance_filter.append({"matric": student.get("matricNumber") or student.get("matric")})
        elif not student:
            attendance_filter.append({"matric": studentId})
            attendance_filter.append({"name": studentId})
        attendance_rows = list(db.attendance.find(
            {"sessionId": {"$in": session_ids}, "$or": attendance_filter},
            {"_id": 0, "sessionId": 1, "userId": 1, "name": 1, "matric": 1,
             "attendanceStatus": 1, "department": 1, "level": 1, "timestamp": 1},
        )) if session_ids else []
        by_session = {row.get("sessionId"): row for row in attendance_rows}
        if not student and not attendance_rows:
            raise HTTPException(status_code=404, detail="Student not found")

        history = []
        student_entry = {
            "userId": (student or {}).get("userId") or (attendance_rows[0].get("userId") if attendance_rows else ""),
            "name": (student or {}).get("name") or (attendance_rows[0].get("name") if attendance_rows else "Unknown student"),
            "matric": (student or {}).get("matricNumber") or (student or {}).get("matric") or (
                attendance_rows[0].get("matric", "") if attendance_rows else ""
            ),
            "department": (student or {}).get("department", ""),
            "level": (student or {}).get("level", ""),
        }
        for session in sessions:
            session_id = session.get("sessionId", "")
            code = session.get("courseCode") or session.get("code") or ""
            attendance = by_session.get(session_id)
            if code not in enrolled_codes and not attendance:
                continue
            attendee = {
                **student_entry,
                "timestamp": attendance.get("timestamp", "") if attendance else "",
                "status": attendance.get("attendanceStatus", "present") if attendance else "absent",
            }
            history.append({
                **_session_summary(session, course_lookup, int(attendee["status"] == "present"), None),
                "attendees": [attendee],
            })
        return {"sessions": history}
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Error loading student attendance history")
        raise HTTPException(status_code=500, detail="Failed to load student attendance history") from e
